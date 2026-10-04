"""
update_ppt_link.py
------------------
Utility script to instantly update the hosted URL on Slide 5 of the presentation.
Usage:
    python update_ppt_link.py "https://your-custom-name.netlify.app"
"""

import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def update_link(new_url: str, ppt_path: str = "Smart_Study_Notes_Generator.pptx"):
    prs = Presentation(ppt_path)
    slide5 = prs.slides[4]  # 5th slide (0-indexed)

    # Search for text containing "PASTE YOUR HOSTED LINK HERE" or ".netlify.app" or ".vercel.app"
    updated = False
    for shape in slide5.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if "PASTE YOUR HOSTED LINK HERE" in p.text or ".netlify.app" in p.text or ".vercel.app" in p.text:
                    p.text = f"LIVE HOSTED APPLICATION LINK:\n{new_url}"
                    p.font.size = Pt(13)
                    p.font.bold = True
                    p.font.color.rgb = RGBColor(11, 15, 25)
                    p.alignment = PP_ALIGN.CENTER
                    updated = True
                    break
        if updated:
            break

    if updated:
        prs.save(ppt_path)
        print(f"Successfully updated Slide 5 with hosted link: {new_url}")
    else:
        print("Could not find the target link placeholder on Slide 5. Please re-run generate_ppt.py.")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python update_ppt_link.py <hosted_url>")
        sys.exit(1)
    update_link(sys.argv[1])
